from __future__ import annotations

import io
import tempfile
import unittest
import zipfile
from pathlib import Path


class TestXlsxSkillCore(unittest.TestCase):
    def setUp(self) -> None:
        import codegaai.core.xlsx_skill as xlsx

        self.xlsx = xlsx
        self.tmp = tempfile.TemporaryDirectory()
        base = Path(self.tmp.name)
        self.old_input = xlsx.XLSX_INPUT_DIR
        self.old_output = xlsx.XLSX_OUTPUT_DIR
        self.old_roots = xlsx.ALLOWED_XLSX_ROOTS
        xlsx.XLSX_INPUT_DIR = base / "temp" / "xlsx"
        xlsx.XLSX_OUTPUT_DIR = base / "outputs" / "xlsx"
        xlsx.ALLOWED_XLSX_ROOTS = (xlsx.XLSX_INPUT_DIR, xlsx.XLSX_OUTPUT_DIR)

    def tearDown(self) -> None:
        self.xlsx.XLSX_INPUT_DIR = self.old_input
        self.xlsx.XLSX_OUTPUT_DIR = self.old_output
        self.xlsx.ALLOWED_XLSX_ROOTS = self.old_roots
        self.tmp.cleanup()

    def test_safe_xlsx_filename_strips_path_traversal(self) -> None:
        self.assertEqual(
            self.xlsx.safe_xlsx_filename("../../finance report.xlsx"),
            "finance_report.xlsx",
        )

    def test_resolve_allowed_path_blocks_outside_workspace(self) -> None:
        outside = Path(self.tmp.name).parent / "outside.xlsx"
        with self.assertRaises(self.xlsx.XlsxSkillError):
            self.xlsx.resolve_allowed_xlsx_path(outside)

    def test_save_uploaded_workbook_rejects_invalid_zip(self) -> None:
        with self.assertRaises(self.xlsx.XlsxSkillError):
            self.xlsx.save_uploaded_workbook("bad.xlsx", b"not an xlsx")

    def test_full_workbook_edit_when_dependencies_exist(self) -> None:
        try:
            import openpyxl  # type: ignore[import-not-found]
            import pandas  # noqa: F401  # type: ignore[import-not-found]
        except ImportError:
            self.skipTest("pandas/openpyxl not installed in this environment")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Month", "Revenue"])
        ws.append(["Jan", 10])
        ws.append(["Feb", 20])
        ws.append(["Mar", 30])
        buf = io.BytesIO()
        wb.save(buf)

        path = self.xlsx.save_uploaded_workbook("sales.xlsx", buf.getvalue())
        summary = self.xlsx.inspect_workbook(path)
        self.assertEqual(summary["sheet_count"], 1)
        self.assertEqual(summary["sheets"][0]["name"], "Sales")

        result = self.xlsx.apply_workbook_operations(
            path,
            [
                {"action": "create_sheet", "sheet": "Report"},
                {"action": "update_cell", "sheet": "Report", "cell": "A1", "value": "Ready"},
                {"action": "format_range", "sheet": "Report", "range": "A1", "fill": "F59E0B", "font": {"bold": True}},
                {"action": "add_chart", "sheet": "Sales", "chart_type": "bar", "data_range": "A1:B4", "anchor": "D2"},
            ],
            "edited.xlsx",
        )
        self.assertTrue(Path(result["output_path"]).exists())
        self.assertIn("created sheet Report", result["applied"])
        self.assertEqual(result["output_url"], "/outputs/xlsx/edited.xlsx")


class TestXlsxApiContract(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        from fastapi.testclient import TestClient
        from codegaai.api.server import app

        cls.client = TestClient(app)

    def test_xlsx_capabilities_endpoint(self) -> None:
        response = self.client.get("/api/files/xlsx/capabilities")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["skill"], "xlsx")
        self.assertIn("add_chart", data["operations"])
        self.assertIn("bar", data["chart_types"])

    def test_xlsx_upload_rejects_invalid_workbook(self) -> None:
        response = self.client.post(
            "/api/files/xlsx/upload",
            files={"file": ("bad.xlsx", io.BytesIO(b"not a workbook"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(response.status_code, 400)

    def test_xlsx_upload_rejects_zip_without_workbook_manifest(self) -> None:
        fake = io.BytesIO()
        with zipfile.ZipFile(fake, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
        fake.seek(0)
        response = self.client.post(
            "/api/files/xlsx/upload",
            files={"file": ("fake.xlsx", fake, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(response.status_code, 400)


if __name__ == "__main__":
    unittest.main()
